"""
Builds the demo files in backend/demo_data/.

  SiteSense_Financials_Template.xlsx          blank template analysts fill in
  Sakura_Precision_Parts_Financials_FY26.xlsx  demo company, analyst's sheet
  Sakura_Precision_Parts_Annual_Report_FY26.pdf demo company, annual report extract
  Kaveri_Logistics_Financials_FY26.xlsx        second demo company (used by seed_demo.py)

Both companies are FICTIONAL. Sakura's numbers contain planted warning signs:
inventory piling up, slower-paying customers, debt doubled, big capex,
92% utilisation claimed while headcount fell, shrinking margin.
One planted disagreement: trade receivables are 91.20 in the Excel sheet
but 93.30 in the annual report, so the app flags a conflict.

Run from the backend folder:  python tools/make_demo_files.py
Needs: pip install reportlab openpyxl
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

OUT = Path(__file__).resolve().parent.parent / "demo_data"
OUT.mkdir(exist_ok=True)

ROWS = [
    # label, code, unit, note
    ("Revenue from operations", "revenue", "Rs crore", "Statement of profit and loss"),
    ("Cost of goods sold", "cogs", "Rs crore", "Statement of profit and loss"),
    ("EBITDA", "ebitda", "Rs crore", "Statement of profit and loss"),
    ("Inventories", "inventory", "Rs crore", "Balance sheet"),
    ("Trade receivables", "receivables", "Rs crore", "Balance sheet"),
    ("Total borrowings", "total_debt", "Rs crore", "Balance sheet, current + non-current"),
    ("Total equity", "equity", "Rs crore", "Balance sheet"),
    ("Capital expenditure", "capex", "Rs crore", "Cash flow, purchase of fixed assets"),
    ("Capacity utilisation (claimed)", "capacity_utilization", "%", "As stated by management"),
    ("Number of employees", "employees", "count", "Permanent employees at year end"),
]

SAKURA = {
    "revenue": (436.70, 412.00), "cogs": (318.00, 298.00), "ebitda": (49.50, 58.00),
    "inventory": (74.10, 49.80), "receivables": (91.20, 62.10), "total_debt": (245.00, 120.00),
    "equity": (322.00, 310.00), "capex": (96.00, 18.00), "capacity_utilization": (92, 78),
    "employees": (1180, 1240),
}
SAKURA_PDF_RECEIVABLES = (93.30, 62.10)

KAVERI = {
    "revenue": (176.00, 182.00), "cogs": (147.00, 150.00), "ebitda": (14.10, 16.40),
    "inventory": (21.00, 20.50), "receivables": (34.50, 30.00), "total_debt": (42.00, 40.00),
    "equity": (88.00, 85.00), "capex": (5.50, 6.00), "capacity_utilization": (72, 70),
    "employees": (405, 410),
}


# ---------------- Excel ----------------

def make_excel(path, company, values):
    wb = Workbook()
    ws = wb.active
    ws.title = "Financials"
    ink = "252A6B"
    thin = Side(style="thin", color="D0D4DD")

    ws["A1"] = "SiteSense financials sheet"
    ws["A1"].font = Font(bold=True, size=14, color=ink)
    ws["A2"] = "Company: " + (company or "")
    ws["A3"] = ("Fill the two year columns. Amounts in Rs crore. "
                "Do not change the Code column, SiteSense reads it.")
    ws["A3"].font = Font(italic=True, color="5E6578")

    header = ["Metric", "Code", "Current year (FY2025-26)", "Previous year (FY2024-25)", "Unit", "Where to find it"]
    ws.append([])
    ws.append(header)
    hrow = ws.max_row
    for col in range(1, len(header) + 1):
        c = ws.cell(row=hrow, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ink)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for label, code, unit, note in ROWS:
        cur, prev = values.get(code, (None, None)) if values else (None, None)
        ws.append([label, code, cur, prev, unit, note])
        r = ws.max_row
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = Border(bottom=thin)
        for col in (3, 4):
            ws.cell(row=r, column=col).number_format = "#,##0.00" if unit != "count" and unit != "%" else "#,##0"
            if values is None:
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FFF8E6")
        ws.cell(row=r, column=2).font = Font(color="8A90A2")

    widths = [32, 22, 24, 24, 10, 38]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[hrow].height = 32
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    wb.save(path)


# ---------------- PDF ----------------

INK = colors.HexColor("#1E3A5F")
GREY = colors.HexColor("#5E6578")
RULE = colors.HexColor("#C9CED8")

h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, textColor=INK, spaceAfter=8, leading=19)
h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, textColor=INK, spaceBefore=10, spaceAfter=6)
body = ParagraphStyle("body", fontName="Helvetica", fontSize=9.5, leading=14, spaceAfter=7)
small = ParagraphStyle("small", fontName="Helvetica", fontSize=8, textColor=GREY, leading=11)


def footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(RULE)
    canvas.line(18 * mm, 15 * mm, A4[0] - 18 * mm, 15 * mm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(GREY)
    canvas.drawString(18 * mm, 10 * mm, "Sakura Precision Parts Limited  |  Annual Report 2025-26")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page " + str(doc.page))
    canvas.drawCentredString(A4[0] / 2, 5.5 * mm,
                             "Fictional company. Demo document prepared for the SiteSense prototype.")
    canvas.restoreState()


def fin_table(title_cols, rows):
    data = [title_cols] + rows
    t = Table(data, colWidths=[86 * mm, 14 * mm, 32 * mm, 32 * mm])
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 9),
        ("TEXTCOLOR", (0, 0), (-1, 0), INK),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, INK),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, RULE),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]
    for i, row in enumerate(rows, start=1):
        if row[0].startswith("Total") or row[0] in ("EBITDA", "Profit after tax"):
            style.append(("FONT", (0, i), (-1, i), "Helvetica-Bold", 9))
    t.setStyle(TableStyle(style))
    return t


def n(v):
    return format(v, ",.2f")


def make_pdf(path):
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=22 * mm,
                            title="Sakura Precision Parts Limited - Annual Report 2025-26 (extract)")
    s = []

    # Cover
    s.append(Spacer(1, 60 * mm))
    s.append(Paragraph("Sakura Precision Parts Limited", ParagraphStyle(
        "cover", fontName="Helvetica-Bold", fontSize=26, textColor=INK, leading=32)))
    s.append(Spacer(1, 6 * mm))
    s.append(Paragraph("Annual Report 2025-26", ParagraphStyle(
        "cover2", fontName="Helvetica", fontSize=16, textColor=GREY, leading=20)))
    s.append(Spacer(1, 4 * mm))
    s.append(Paragraph("Extract: management discussion and financial statements", body))
    s.append(Spacer(1, 80 * mm))
    s.append(Paragraph("Precision machined components for automotive and industrial customers. "
                       "Plants at Pune and Chakan, Maharashtra.", small))
    s.append(PageBreak())

    # Management discussion (claims that Phase 2 will compare with the site)
    s.append(Paragraph("Management discussion and analysis", h1))
    s.append(Paragraph(
        "FY2025-26 was a year of investment for Sakura. Revenue grew 6% to Rs 436.7 crore, "
        "supported by steady demand from our two largest automotive customers.", body))
    s.append(Paragraph(
        "Our plants operated at a capacity utilisation of 92%, up from 78% in the previous year. "
        "We now run close to full capacity across all three shifts at both plants.", body))
    s.append(Paragraph(
        "During the year we commissioned a new five-axis CNC machining line at the Pune plant. "
        "The line is fully operational and is already contributing to output. "
        "The investment was funded through a term loan from our banking partners.", body))
    s.append(Paragraph(
        "Inventory levels are healthy and in line with expected demand. We have built "
        "finished-goods stock ahead of new programmes scheduled to start in the coming year.", body))
    s.append(Paragraph(
        "Margins were affected by higher steel prices during the second half of the year. "
        "We expect them to recover as price revisions with customers take effect.", body))
    s.append(Paragraph("Outlook", h2))
    s.append(Paragraph(
        "With the new line in place, we are confident of strong growth in FY2026-27. "
        "Our order book remains robust and our customer relationships are long-standing.", body))
    s.append(PageBreak())

    # Profit and loss
    s.append(Paragraph("Statement of profit and loss for the year ended 31 March 2026", h1))
    s.append(Paragraph("All amounts in Rs crore unless stated otherwise", small))
    s.append(Spacer(1, 4 * mm))
    head = ["Particulars", "Note", "FY 2025-26", "FY 2024-25"]
    s.append(fin_table(head, [
        ["Revenue from operations", "18", n(436.70), n(412.00)],
        ["Other income", "19", n(3.10), n(2.40)],
        ["Cost of goods sold", "20", n(318.00), n(298.00)],
        ["Employee benefits expense", "21", n(41.20), n(39.80)],
        ["Other expenses", "22", n(28.00), n(16.20)],
        ["EBITDA", "", n(49.50), n(58.00)],
        ["Depreciation and amortisation", "23", n(19.40), n(14.10)],
        ["Finance costs", "24", n(18.60), n(9.20)],
        ["Profit before tax", "", n(14.60), n(37.10)],
        ["Tax expense", "25", n(3.70), n(9.40)],
        ["Profit after tax", "", n(10.90), n(27.70)],
    ]))
    s.append(PageBreak())

    # Balance sheet
    s.append(Paragraph("Balance sheet as at 31 March 2026", h1))
    s.append(Paragraph("All amounts in Rs crore", small))
    s.append(Spacer(1, 4 * mm))
    s.append(fin_table(head, [
        ["Property, plant and equipment", "3", n(298.40), n(221.80)],
        ["Inventories", "7", n(74.10), n(49.80)],
        ["Trade receivables", "8", n(SAKURA_PDF_RECEIVABLES[0]), n(SAKURA_PDF_RECEIVABLES[1])],
        ["Cash and cash equivalents", "9", n(12.30), n(21.60)],
        ["Other assets", "10", n(208.20), n(193.70)],
        ["Total assets", "", n(686.30), n(549.00)],
        ["Equity share capital", "11", n(50.00), n(50.00)],
        ["Other equity", "12", n(272.00), n(260.00)],
        ["Total equity", "", n(322.00), n(310.00)],
        ["Total borrowings", "13", n(245.00), n(120.00)],
        ["Trade payables", "14", n(86.30), n(79.10)],
        ["Other liabilities", "15", n(33.00), n(39.90)],
        ["Total equity and liabilities", "", n(686.30), n(549.00)],
    ]))
    s.append(Spacer(1, 8 * mm))

    s.append(Paragraph("Other information", h2))
    s.append(fin_table(["Item", "", "FY 2025-26", "FY 2024-25"], [
        ["Capital expenditure", "", n(96.00), n(18.00)],
        ["Number of employees (nos.)", "", "1,180", "1,240"],
        ["Plants (nos.)", "", "2", "2"],
    ]))

    doc.build(s, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    make_excel(OUT / "SiteSense_Financials_Template.xlsx", "", None)
    make_excel(OUT / "Sakura_Precision_Parts_Financials_FY26.xlsx", "Sakura Precision Parts Limited", SAKURA)
    make_excel(OUT / "Kaveri_Logistics_Financials_FY26.xlsx", "Kaveri Logistics Private Limited", KAVERI)
    make_pdf(OUT / "Sakura_Precision_Parts_Annual_Report_FY26.pdf")
    for p in sorted(OUT.iterdir()):
        print("wrote", p.name)
