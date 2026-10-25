"""
Reading numbers out of the uploaded documents.

Order of trust (highest first):
  1. Excel template     - read cell by cell, exact
  2. PDF, rule-based    - finds rows like "Inventories 74.10 49.80" by their label
  3. AI (optional)      - only for numbers still missing, and only kept if the
                          number really appears on the page the AI cites

If two sources disagree by more than 1%, the metric is flagged as a conflict
and the analyst decides.
"""

import io
import json
import re

import openpyxl
import pdfplumber
import requests
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel

import finance
import settings
import storage
from auth import get_current_user, require_analyst
from docs import read_document

router = APIRouter(prefix="/api/visits", tags=["numbers"])

PRIORITY = {"excel": 3, "pdf": 2, "ai": 1, "excel_ai": 1}
UNIT_TO_CRORE = {"crore": 1.0, "lakh": 0.01, "million": 0.1, "rupees": 1e-7}


def to_crore(found, unit):
    """Convert money figures to Rs crore using the unit the analyst said the file uses."""
    factor = UNIT_TO_CRORE.get(unit, 1.0)
    if factor == 1.0:
        return
    for key, v in found.items():
        if finance.UNITS[key] == "cr":
            for field in ("current", "previous"):
                if v.get(field) is not None:
                    v[field] = round(v[field] * factor, 4)
NUM_RE = re.compile(r"\(?-?\d[\d,]*(?:\.\d+)?\)?")


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("Rs", "").replace("%", "").strip()
    if text in ("", "-", "--", "na", "n/a"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def _match_label(text, table_row_only=False):
    """Which metric does this row label belong to? Longest synonym wins.

    table_row_only: the rest of the line must be numbers only (plus a few letters
    like "Rs"), so a sentence such as "Revenue grew 6% to Rs 436.7 crore" is ignored.
    """
    low = re.sub(r"\s+", " ", str(text).lower()).strip()
    # Drop list numbering like "1. ", "(a) ", "iv) " at the start of a row
    low = re.sub(r"^\(?[0-9a-z]{1,4}[.)]\s+", "", low)
    best = None
    for key in finance.METRIC_KEYS:
        for syn in finance.SYNONYMS[key]:
            if low.startswith(syn) and (best is None or len(syn) > best[1]):
                best = (key, len(syn))
    if best is None:
        return None
    if table_row_only:
        leftover_letters = re.sub(r"[^a-z]", "", low[best[1]:])
        if len(leftover_letters) > 4:
            return None
    return best[0]


# ---------- Excel ----------

def read_excel(content, filename):
    found, log = {}, []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return found, [filename + ": could not open as an Excel file."]

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_row, code_col, cur_col, prev_col = None, None, None, None
        for r_idx, row in enumerate(rows[:30]):
            cells = [str(c).strip().lower() if c is not None else "" for c in row]
            if "code" in cells:
                header_row, code_col = r_idx, cells.index("code")
                for c_idx, c in enumerate(cells):
                    if "current" in c and cur_col is None:
                        cur_col = c_idx
                    if "previous" in c and prev_col is None:
                        prev_col = c_idx
                break

        for r_idx, row in enumerate(rows):
            if header_row is not None and r_idx <= header_row:
                continue
            key, cur, prev = None, None, None
            if header_row is not None and cur_col is not None and prev_col is not None:
                code = str(row[code_col]).strip().lower() if code_col < len(row) and row[code_col] else ""
                if code in finance.METRIC_KEYS:
                    key = code
                    cur = to_number(row[cur_col]) if cur_col < len(row) else None
                    prev = to_number(row[prev_col]) if prev_col < len(row) else None
            else:
                # No template header: first text cell is the label, next two numbers are the years.
                label = next((c for c in row if isinstance(c, str) and c.strip()), None)
                if label:
                    key = _match_label(label)
                    nums = [n for n in (to_number(c) for c in row) if n is not None]
                    if len(nums) >= 2:
                        cur, prev = nums[0], nums[1]
            if key and key not in found and (cur is not None or prev is not None):
                found[key] = {
                    "current": cur,
                    "previous": prev,
                    "method": "excel",
                    "source": filename + ", sheet " + ws.title + ", row " + str(r_idx + 1),
                }
    log.append(filename + ": read " + str(len(found)) + " numbers from the spreadsheet.")
    return found, log


def excel_ai_fill(missing, content, filename):
    """For spreadsheets in any layout. The AI only POINTS at cells (sheet and two cell addresses).
    The numbers are then read from those cells by our own code, so the AI never supplies a number."""
    api_key = settings.get("LLM_API_KEY")
    if not missing or not api_key:
        return {}, []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        return {}, []

    lines, total = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=300):
            cells = [c for c in row if getattr(c, "value", None) not in (None, "")]
            if not cells:
                continue
            line = ws.title + "!" + " | ".join(c.coordinate + "=" + str(c.value)[:60] for c in cells[:8])
            if total + len(line) > 40000:
                break
            lines.append(line)
            total += len(line)
    wanted = {k: finance.LABELS[k] for k in missing}
    prompt = (
        "Below are the non-empty rows of a company's financial spreadsheet, as SHEET!CELL=value.\n"
        "For each metric, find the row that holds it. Give the cell of the CURRENT year and the cell of "
        "the PREVIOUS year in that same row. Do not copy or calculate numbers; give cell addresses only.\n"
        "Metrics: " + json.dumps(wanted) + "\n"
        'Return JSON only: {"metrics": {"<key>": {"sheet": "<sheet>", "current_cell": "C5", "previous_cell": "D5"}}}\n'
        "If a metric is not clearly there, leave it out.\n\n" + "\n".join(lines)
    )
    try:
        res = requests.post(
            settings.get("LLM_BASE_URL").rstrip("/") + "/chat/completions",
            headers={"Authorization": "Bearer " + api_key},
            json={"model": settings.get("LLM_MODEL"), "temperature": 0,
                  "response_format": {"type": "json_object"},
                  "messages": [{"role": "user", "content": prompt}]},
            timeout=60,
        )
        res.raise_for_status()
        answer = json.loads(res.json()["choices"][0]["message"]["content"])
    except Exception as e:
        return {}, [filename + ": AI could not read this spreadsheet (" + type(e).__name__ + ")."]

    found, log = {}, []
    for key, item in (answer.get("metrics") or {}).items():
        if key not in missing or not isinstance(item, dict) or item.get("sheet") not in wb.sheetnames:
            continue
        ws = wb[item["sheet"]]
        try:
            cur_cell, prev_cell = ws[str(item.get("current_cell"))], ws[str(item.get("previous_cell"))]
            cur, prev = to_number(cur_cell.value), to_number(prev_cell.value)
        except Exception:
            continue
        # Our checks: both cells hold numbers and sit in the same row.
        if cur is None or prev is None or cur_cell.row != prev_cell.row:
            log.append("AI pointed at cells for " + finance.LABELS[key] + " that did not hold two numbers in one row, so it was rejected.")
            continue
        found[key] = {"current": cur, "previous": prev, "method": "excel_ai",
                      "source": filename + ", sheet " + item["sheet"] + ", row " + str(cur_cell.row) + " (row chosen by AI)"}
    log.append(filename + ": AI found " + str(len(found)) + " of " + str(len(missing)) + " missing numbers (read from the cells by code).")
    return found, log


# ---------- PDF ----------

UTIL_SENTENCE = re.compile(
    r"capacity utili[sz]ation[^%]{0,80}?(\d{1,3}(?:\.\d+)?)\s*%[^%]{0,60}?(\d{1,3}(?:\.\d+)?)\s*%",
    re.IGNORECASE,
)


def read_pdf(content, filename):
    found, log, pages = {}, [], []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                pages.append(text)
                for line in text.splitlines():
                    key = _match_label(line, table_row_only=True)
                    if not key or key in found:
                        continue
                    numbers = [to_number(n) for n in NUM_RE.findall(line)]
                    numbers = [n for n in numbers if n is not None]
                    if len(numbers) >= 2:
                        cur, prev = numbers[-2], numbers[-1]
                        if key == "capex":
                            # Cash flow statements print spending as (96.00)
                            cur, prev = abs(cur), abs(prev)
                        found[key] = {
                            "current": cur,
                            "previous": prev,
                            "method": "pdf",
                            "source": filename + ", page " + str(page_no),
                        }
                if "capacity_utilization" not in found:
                    m = UTIL_SENTENCE.search(text.replace("\n", " "))
                    if m:
                        found["capacity_utilization"] = {
                            "current": float(m.group(1)),
                            "previous": float(m.group(2)),
                            "method": "pdf",
                            "source": filename + ", page " + str(page_no),
                        }
    except Exception:
        return found, [filename + ": could not read this PDF. Is it scanned or password-protected?"], pages

    if not any(p.strip() for p in pages):
        log.append(filename + ": no text found. It may be a scanned PDF, which is not supported yet.")
    else:
        log.append(filename + ": found " + str(len(found)) + " numbers by matching row labels.")
    return found, log, pages


# ---------- Optional AI fill-in ----------

def select_pages(pdf_pages, missing, max_pages):
    """Score every page by how many of the missing metrics' label words it contains, and keep the
    best few. Plain keyword counting: no embeddings, no vector database. Returns
    [(filename, page_no, text)] in reading order."""
    words = [w.lower() for k in missing for w in finance.SYNONYMS.get(k, [])]
    scored = []
    for filename, pages in pdf_pages:
        for i, text in enumerate(pages, start=1):
            low = text.lower()
            hits = sum(low.count(w) for w in words)
            numbers = len(NUM_RE.findall(text))
            if hits and numbers >= 4:  # a page of prose with a keyword is not a table
                scored.append((hits, filename, i, text))
    scored.sort(key=lambda s: s[0], reverse=True)
    best = sorted(scored[:max_pages], key=lambda s: (s[1], s[2]))
    return [(f, i, t) for _, f, i, t in best]


def ai_fill(missing, pdf_pages):
    """pdf_pages: list of (filename, [page texts]). Returns (found, log)."""
    api_key = settings.get("LLM_API_KEY")
    if not missing:
        return {}, []
    if not api_key:
        return {}, ["AI reading is off (no LLM_API_KEY), so missing numbers were left for you to fill in."]
    if not pdf_pages:
        return {}, []

    selected = select_pages(pdf_pages, missing, settings.get_int("LLM_MAX_PAGES", 6))
    if not selected:
        return {}, ["No page of the PDFs looked like it holds the missing numbers, so AI was not used."]
    chunks, total = [], 0
    for filename, i, text in selected:
        block = "=== FILE: " + filename + " | PAGE: " + str(i) + " ===\n" + text[:12000] + "\n"
        if total + len(block) > 60000:
            break
        chunks.append(block)
        total += len(block)

    wanted = {k: finance.LABELS[k] for k in missing}
    prompt = (
        "You read company financial reports. Find these metrics for the CURRENT and "
        "PREVIOUS financial year. Give amounts exactly as printed; do not convert units. "
        "Capacity utilisation is a percent. Employees is a count.\n"
        "Metrics: " + json.dumps(wanted) + "\n"
        "Return JSON only, in this exact shape:\n"
        '{"metrics": {"<key>": {"current": number or null, "previous": number or null, '
        '"file": "<file name>", "page": <page number>}}}\n'
        "If a metric is not clearly stated, return null for it. Do not calculate or guess.\n\n"
        + "".join(chunks)
    )
    base_url = settings.get("LLM_BASE_URL").rstrip("/")
    try:
        res = requests.post(
            base_url + "/chat/completions",
            headers={"Authorization": "Bearer " + api_key},
            json={
                "model": settings.get("LLM_MODEL"),
                "temperature": 0,
                "response_format": {"type": "json_object"},
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=60,
        )
        res.raise_for_status()
        answer = json.loads(res.json()["choices"][0]["message"]["content"])
    except Exception as e:
        return {}, ["AI reading failed (" + type(e).__name__ + "). Missing numbers were left for you to fill in."]

    page_lookup = {(f, i): t for f, pages in pdf_pages for i, t in enumerate(pages, start=1)}
    found, log = {}, []
    for key, item in (answer.get("metrics") or {}).items():
        if key not in missing or not isinstance(item, dict):
            continue
        cur, prev = to_number(item.get("current")), to_number(item.get("previous"))
        if cur is None and prev is None:
            continue
        try:
            page_no = int(item.get("page"))
        except (TypeError, ValueError):
            page_no = -1
        text = page_lookup.get((item.get("file"), page_no))
        # Anti-hallucination check: the numbers must be printed on the cited page.
        printed = [to_number(n) for n in NUM_RE.findall(text or "")]
        ok = text is not None and all(
            v is None or any(p is not None and abs(p - v) < 0.006 for p in printed) for v in (cur, prev)
        )
        if ok:
            found[key] = {
                "current": cur, "previous": prev, "method": "ai",
                "source": item.get("file") + ", page " + str(page_no),
            }
        else:
            log.append("AI suggested a value for " + finance.LABELS[key] +
                       " but it was not printed on the page it cited, so it was rejected.")
    log.append("AI filled " + str(len(found)) + " of " + str(len(missing)) + " missing numbers (each checked against its page).")
    return found, log


# ---------- Merge ----------

def _differs(a, b):
    if a is None or b is None:
        return False
    return abs(a - b) > max(abs(a), abs(b)) * 0.01


def run_extraction(visit):
    if not visit.get("docs"):
        raise HTTPException(400, "Upload at least one document first.")

    candidates = {k: [] for k in finance.METRIC_KEYS}
    log, pdf_pages, excel_docs, unit_of = [], [], [], {}
    for meta in visit["docs"]:
        _, content = read_document(visit, meta["doc_id"])
        unit_of[meta["filename"]] = meta.get("unit", "crore")
        if meta["ext"] == ".xlsx":
            found, msgs = read_excel(content, meta["filename"])
            excel_docs.append((meta["filename"], content))
        else:
            found, msgs, pages = read_pdf(content, meta["filename"])
            pdf_pages.append((meta["filename"], pages))
        to_crore(found, unit_of[meta["filename"]])
        log.extend(msgs)
        for key, value in found.items():
            candidates[key].append(value)

    # Still missing? Step 1: AI points at cells in the spreadsheets. Step 2: AI reads the best PDF pages.
    for filename, content in excel_docs:
        missing = [k for k in finance.METRIC_KEYS if not candidates[k]]
        ai_found, ai_log = excel_ai_fill(missing, content, filename)
        log.extend(ai_log)
        to_crore(ai_found, unit_of[filename])
        for key, value in ai_found.items():
            candidates[key].append(value)
    missing = [k for k in finance.METRIC_KEYS if not candidates[k]]
    ai_found, ai_log = ai_fill(missing, pdf_pages)
    log.extend(ai_log)
    for key, value in ai_found.items():
        to_crore({key: value}, unit_of.get(value["source"].rsplit(", page", 1)[0], "crore"))
        candidates[key].append(value)

    metrics = {}
    for key in finance.METRIC_KEYS:
        options = sorted(candidates[key], key=lambda c: PRIORITY[c["method"]], reverse=True)
        chosen = options[0] if options else {}
        conflict = any(
            _differs(chosen.get("current"), o["current"]) or _differs(chosen.get("previous"), o["previous"])
            for o in options[1:]
        )
        metrics[key] = {
            "label": finance.LABELS[key],
            "unit": finance.UNITS[key],
            "current": chosen.get("current"),
            "previous": chosen.get("previous"),
            "method": chosen.get("method"),
            "source": chosen.get("source"),
            "conflict": conflict,
            "candidates": options,
        }
    raw = {
        "extracted_at": storage.now_iso(),
        "documents": [d["doc_id"] for d in visit["docs"]],
        "metrics": metrics,
        "missing": [k for k in finance.METRIC_KEYS if metrics[k]["current"] is None],
        "conflicts": [k for k in finance.METRIC_KEYS if metrics[k]["conflict"]],
        "log": log,
    }
    storage.write_json(storage.visit_path(visit["visit_id"], "metrics_raw.json"), raw)
    return raw


def confirm_metrics(visit, user, values):
    """values: {key: {"current": x, "previous": y, "source": optional}}"""
    storage.ensure_not_locked(visit)
    vid = visit["visit_id"]
    raw = storage.read_json(storage.visit_path(vid, "metrics_raw.json"))
    if raw is None:
        raise HTTPException(400, "Read the numbers from the documents first.")

    confirmed = {}
    for key in finance.METRIC_KEYS:
        r = raw["metrics"][key]
        v = values.get(key, {})
        cur, prev = to_number(v.get("current")), to_number(v.get("previous"))
        edited = cur != r["current"] or prev != r["previous"]
        source, method = r["source"], r["method"]
        if edited:
            # Did the analyst pick another document's value, or type their own?
            match = next((c for c in r["candidates"] if c["current"] == cur and c["previous"] == prev), None)
            source = match["source"] if match else "Entered by analyst"
            method = match["method"] if match else "analyst"
            storage.append_audit(vid, user["email"], "number_edited", {
                "metric": key, "old": [r["current"], r["previous"]], "new": [cur, prev], "source": source,
            })
        confirmed[key] = {
            "label": r["label"], "unit": r["unit"], "current": cur, "previous": prev,
            "source": source, "method": method, "edited": edited,
        }

    doc = {
        "confirmed_at": storage.now_iso(),
        "confirmed_by": user["email"],
        "metrics": confirmed,
        "derived": finance.compute_all(confirmed),
    }
    storage.write_json(storage.visit_path(vid, "metrics.json"), doc)
    storage.append_audit(vid, user["email"], "numbers_confirmed")
    return doc


# ---------- routes ----------

@router.post("/{visit_id}/extract")
def extract(visit_id: str, user=Depends(get_current_user)):
    require_analyst(user)
    visit = storage.load_visit_for_user(visit_id, user)
    storage.ensure_not_locked(visit)
    raw = run_extraction(visit)
    storage.append_audit(visit_id, user["email"], "numbers_read", {
        "found": len(finance.METRIC_KEYS) - len(raw["missing"]), "conflicts": len(raw["conflicts"]),
    })
    return raw


@router.get("/{visit_id}/metrics")
def get_metrics(visit_id: str, user=Depends(get_current_user)):
    storage.load_visit_for_user(visit_id, user, read_only=True)
    return {
        "raw": storage.read_json(storage.visit_path(visit_id, "metrics_raw.json")),
        "confirmed": storage.read_json(storage.visit_path(visit_id, "metrics.json")),
    }


class MetricsIn(BaseModel):
    metrics: dict[str, dict]


@router.put("/{visit_id}/metrics")
def put_metrics(visit_id: str, body: MetricsIn, user=Depends(get_current_user)):
    require_analyst(user)
    visit = storage.load_visit_for_user(visit_id, user)
    return confirm_metrics(visit, user, body.metrics)
